"""Exporte data/films_du_monde.json en films-du-monde.xlsx."""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
rows = json.loads((ROOT / "data" / "films_du_monde.json").read_text())

COLS = [
    ("Pays", "pays", 22), ("Continent", "continent", 13), ("Niveau", "niveau", 8),
    ("Fiabilité", "fiabilite", 22), ("Titre original", "titre_original", 34),
    ("Titre français", "title_fr", 32), ("Année", "annee", 7), ("Langue", "langue", 8),
    ("Animation", "animation", 10), ("Durée (min)", "runtime", 10),
    ("Critère", "critere", 34), ("Où le voir en France", "where_fr", 48),
    ("Source du classement", "source", 40), ("Fiche TMDB", "tmdb_url", 36),
    ("2e du classement", "second", 28), ("Le plus voté sur TMDB", "tmdb_top", 34),
    ("Choix critique (Films Around the World)", "fatw", 34),
    ("Phénomène vidéo / YouTube", "video", 40), ("Notes", "note", 60),
]
FILL = {"A": "D9EAD3", "B": "FFF2CC", "C": "FCE5CD", "D": "F4CCCC", "-": "EEEEEE"}
FONT = "Arial"

wb = Workbook()
ws = wb.active
ws.title = "Films"
ws.append([c[0] for c in COLS])
for r in rows:
    vals = []
    for _, key, _ in COLS:
        v = r.get(key)
        if key == "animation":
            v = "oui" if v else ""
        if key == "annee" and v:
            v = str(v)
        vals.append(v if v not in (None, False) else "")
    ws.append(vals)

for i, (_, key, width) in enumerate(COLS, 1):
    ws.column_dimensions[get_column_letter(i)].width = width
for row in ws.iter_rows(min_row=1):
    for cell in row:
        cell.font = Font(name=FONT, size=10, bold=cell.row == 1)
        cell.alignment = Alignment(vertical="top", wrap_text=cell.column > 10)
lvl_col = [k for _, k, _ in COLS].index("niveau") + 1
for row in ws.iter_rows(min_row=2):
    lvl = row[lvl_col - 1].value or "-"
    row[lvl_col - 1].fill = PatternFill("solid", fgColor=FILL.get(lvl, "FFFFFF"))
for key in ("source", "tmdb_url"):
    c = [k for _, k, _ in COLS].index(key) + 1
    for row in ws.iter_rows(min_row=2, min_col=c, max_col=c):
        if row[0].value and str(row[0].value).startswith("http"):
            row[0].hyperlink = row[0].value
            row[0].font = Font(name=FONT, size=10, color="1155CC", underline="single")
ref = f"A1:{get_column_letter(len(COLS))}{len(rows) + 1}"
tab = Table(displayName="Films", ref=ref)
tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
ws.add_table(tab)
ws.freeze_panes = "B2"

# Résumé
s = wb.create_sheet("Résumé")
s.append(["Niveau", "Signification", "Pays"])
legend = [("A", "Box-office national sourcé : le film national le plus vu dans le pays"),
          ("B", "Pas de box-office : film national le plus voté sur TMDB (50 votes ou plus)"),
          ("C", "Données trop rares : choix éditorial « native » de Films Around the World"),
          ("D", "Dernier recours : film national le plus voté sur TMDB, moins de 50 votes"),
          ("-", "Aucun film trouvé")]
for i, (lvl, txt) in enumerate(legend, 2):
    s.append([lvl, txt, f'=COUNTIF(Films!$C$2:$C${len(rows) + 1},"{lvl}")'])
    s.cell(i, 1).fill = PatternFill("solid", fgColor=FILL[lvl])
s.append(["Total", "", f"=SUM(C2:C{len(legend) + 1})"])
s.append([])
s.append(["Disponibles en France (streaming, location ou achat)", "",
          f'=COUNTIF(Films!$L$2:$L${len(rows) + 1},"?*")'])
s.column_dimensions["A"].width = 12
s.column_dimensions["B"].width = 80
s.column_dimensions["C"].width = 10
for row in s.iter_rows():
    for cell in row:
        cell.font = Font(name=FONT, size=10, bold=cell.row == 1 or cell.value == "Total")

# Méthode
m = wb.create_sheet("Méthode")
for line in (ROOT / "docs" / "methode.md").read_text().splitlines():
    m.append([line])
    m.cell(m.max_row, 1).font = Font(name=FONT, size=10, bold=line.startswith("#"))
m.column_dimensions["A"].width = 120

out = ROOT / "films-du-monde.xlsx"
wb.save(out)
print(out)
